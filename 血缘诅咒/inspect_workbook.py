from pathlib import Path
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
skill=Path('C:/Users/DHXM/.codex/plugins/cache/openai-primary-runtime/spreadsheets/26.1007.11041/skills/spreadsheets')
print((skill/'SKILL.md').read_text(encoding='utf-8'))
print('API FILES', '\n'.join(str(p) for p in skill.rglob('*') if p.is_file()))
print('FILES')
for p in Path('.').rglob('*'):
    if p.suffix.lower() in ['.xlsx','.png','.jpg','.jpeg','.webp']:
        if not any(part.startswith('.') for part in p.parts): print(p)
import openpyxl
w=openpyxl.load_workbook('场景怪物表.xlsx')
for s in w:
    print('SHEET',s.title,s.max_row,s.max_column)
    for row in s:
        vals=['%s=%s'%(c.coordinate,c.value) for c in row if c.value is not None and (s.title!='怪物表' or c.column==1)]
        if vals: print(' | '.join(vals))
