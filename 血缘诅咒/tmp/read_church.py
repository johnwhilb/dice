from pathlib import Path
import openpyxl,sys
sys.stdout.reconfigure(encoding='utf-8')
p=Path.cwd()/'\u6218\u6597\u573a\u666f'/'\u573a\u666f\u602a\u7269\u6b66\u5668\u8868.xlsx'
w=openpyxl.load_workbook(p,data_only=True)
for s in w:
 print('SHEET',s.title,s.max_row)
 if '\u602a\u7269' not in s.title: continue
 for row in s.iter_rows(values_only=True):
  if row[0] and any(k in str(row[0]) for k in ['\u6559\u4f1a\u4ec6\u4ece','\u6559\u4f1a\u5de8\u4eba']): print(repr(row))
