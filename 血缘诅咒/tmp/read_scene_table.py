from pathlib import Path
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')
root=Path.cwd()
p=root / '\u6218\u6597\u573a\u666f' / '\u573a\u666f\u602a\u7269\u6b66\u5668\u8868.xlsx'
w=openpyxl.load_workbook(p,data_only=True)
for s in w:
 print('\nSHEET',s.title,s.max_row,s.max_column)
 for row in s.iter_rows(values_only=True):
  if any(v is not None for v in row): print(repr(row))
