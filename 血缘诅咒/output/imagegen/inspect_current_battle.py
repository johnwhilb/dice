from pathlib import Path
import sys, json
import openpyxl
from PIL import Image
sys.stdout.reconfigure(encoding='utf-8')
root=Path(__file__).resolve().parents[2]
print('REFERENCES')
for folder in ['画风锚点','可用界面','战斗场景']:
    for p in (root/folder).rglob('*'):
        if p.suffix.lower() in ['.png','.jpg','.jpeg','.webp']:
            with Image.open(p) as im: print(str(p),im.size)
print('WORKBOOK')
w=openpyxl.load_workbook(root/'战斗场景/场景怪物武器表.xlsx',data_only=True)
for s in w:
    print('SHEET',s.title,s.max_row,s.max_column)
    for row in s:
        wanted = s.title=='场景表' and row[0].row<=2 or s.title=='怪物表' and row[0].row<=5 or s.title=='武器表' and row[0].row<=3
        cells=[f'{c.coordinate}: {c.value}' for c in row if c.value is not None and (wanted or c.column==1)]
        if cells: print(' | '.join(cells))
print('DRAWING MEDIA')
for s in w:
    for i,im in enumerate(s._images):
        a=im.anchor
        if s.title=='场景表' and a._from.row==1 or s.title=='怪物表' and a._from.row in (1,2,3):
            p=root/'output/imagegen'/f'current-{s.title}-{a._from.row+1}.png'
            p.write_bytes(im._data())
            print(s.title,i,'row',a._from.row+1,'col',a._from.col+1,'file',p)
